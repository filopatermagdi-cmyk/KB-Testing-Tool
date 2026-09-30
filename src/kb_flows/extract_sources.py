#!/usr/bin/env python3
"""
extract_sources.py — HK "Knowledge Hub" 66-docx batch extractor (AR + EN).

Reads every .docx under an AR folder and an EN folder, extracts each one's
structured text (paragraphs + table rows, same approach as extract_chunks.py),
and pairs them by index into one sources.json that later tools turn into
Claude prompts + per-flow .flow.json files.

Usage:
    python src/kb_flows/extract_sources.py \
        --ar "<KB (Arabic) ... Knowledge Hub>" \
        --en "<KB (English) ... Knowledge Hub>" \
        --out data/kb_flows/sources.json

Optional:
    --pairing data/kb_flows/pairing.json   # [{ar,en}] manual overrides

Output: data/kb_flows/sources.json
"""
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

ARABIC_RE = re.compile(r"[\u0600-\u06FF]")

SPLIT_CHARS = re.compile(r"[\s\-_.]+")


def slugify(name: str) -> str:
    base = name or ""
    base = re.sub(r"\([^)]*\)", "", base).strip()
    parts = [p for p in SPLIT_CHARS.split(base) if p]
    return "_".join(parts).upper() or "FLOW"


def detect_language(text: str) -> str:
    return "ar" if ARABIC_RE.search(text) else "en"


def extract_docx_text(docx_path: str) -> dict:
    """Structured text of one docx: paragraphs + table rows (deduped)."""
    doc = Document(docx_path)
    paragraphs = []
    seen = set()
    for p in doc.paragraphs:
        text = p.text.strip()
        if text and text not in seen:
            seen.add(text)
            paragraphs.append(text)
    tables = []
    seen_rows = set()
    for t_idx, table in enumerate(doc.tables):
        row_texts = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            rtext = " | ".join(cells)
            if rtext and rtext not in seen_rows:
                seen_rows.add(rtext)
                row_texts.append(rtext)
        if row_texts:
            tables.append({"table": t_idx + 1, "rows": row_texts})
    full = "\n".join(paragraphs + [r for t in tables for r in t["rows"]])
    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "full_text": full,
        "language": detect_language(full) if full else "?",
    }


def read_match(pairing_path: str) -> list[dict]:
    if not pairing_path:
        return []
    p = Path(pairing_path)
    if not p.exists():
        return []
    return json.loads(p.read_text(encoding="utf-8"))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ar", required=True, help="Folder with Arabic .docx")
    ap.add_argument("--en", required=True, help="Folder with English .docx")
    ap.add_argument("--out", required=True, help="Output sources.json")
    ap.add_argument("--pairing", default="", help="Optional manual pairing.json")
    args = ap.parse_args()

    ar_dir = Path(args.ar)
    en_dir = Path(args.en)
    if not ar_dir.is_dir() or not en_dir.is_dir():
        print(f"Invalid dir(s): {ar_dir} | {en_dir}", file=sys.stderr)
        raise SystemExit(1)

    ar_files = sorted(ar_dir.glob("*.docx"))
    en_files = sorted(en_dir.glob("*.docx"))
    manual = read_match(args.pairing)
    manual_map = {m["ar"].strip().lower(): m["en"].strip().lower() for m in manual}

    ar_extracted = []
    for f in ar_files:
        data = extract_docx_text(str(f))
        ar_extracted.append(
            {"file": f.name, "slug": slugify(f.stem), **data}
        )
    en_extracted = []
    for f in en_files:
        data = extract_docx_text(str(f))
        en_extracted.append(
            {"file": f.name, "slug": slugify(f.stem), **data}
        )

    by_name = {e["file"].strip().lower(): e for e in en_extracted}
    pairs = []
    used_en = set()
    ar_unpaired = []
    for i, a in enumerate(ar_extracted):
        manual_en = manual_map.get(a["file"].strip().lower())
        if manual_en and manual_en in by_name and manual_en not in used_en:
            en = by_name[manual_en]
            used_en.add(manual_en)
            pairs.append({"index": i + 1, "ar": a, "en": en})
            continue
        if i < len(en_extracted) and en_extracted[i]["file"].strip().lower() not in used_en:
            en = en_extracted[i]
            used_en.add(en["file"].strip().lower())
            pairs.append({"index": i + 1, "ar": a, "en": en})
            continue
        ar_unpaired.append(a)
    en_unpaired = [e for e in en_extracted if e["file"].strip().lower() not in used_en]

    out = {
        "version": 1,
        "count": len(pairs),
        "pairs": pairs,
        "unpaired_ar": [e["file"] for e in ar_unpaired],
        "unpaired_en": [e["file"] for e in en_unpaired],
        "ar_total": len(ar_extracted),
        "en_total": len(en_extracted),
    }
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Extracted {len(pairs)} pairs (AR {len(ar_extracted)} / EN {len(en_extracted)}) -> {args.out}")
    if out["unpaired_ar"]:
        print(f"  unpaired AR: {len(out['unpaired_ar'])} -> {', '.join(out['unpaired_ar'])}")
    if out["unpaired_en"]:
        print(f"  unpaired EN: {len(out['unpaired_en'])} -> {', '.join(out['unpaired_en'])}")
    for p in pairs[:5]:
        ar_text = p["ar"]["full_text"][:60].replace("\n", " ")
        en_text = p["en"]["full_text"][:60].replace("\n", " ")
        print(f"  [{p['index']}] {p['ar']['file']}  <->  {p['en']['file']}\n      AR: {ar_text}\n      EN: {en_text}")


if __name__ == "__main__":
    main()